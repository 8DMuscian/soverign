"""Graph nodes for the Sovereign AI orchestrator.

Each node function receives the full graph state and returns a partial
state dict with the keys it updated.  LangGraph merges the returned
dict into the running state automatically.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import docker
from langchain_core.messages import AIMessage, HumanMessage, SystemMessage
from langgraph.runtime import Runtime

from .state import OrchestratorState

# ────────────────────────────────────────────────────────────
# CONTEXT SCHEMA (passed at invoke time)
# ────────────────────────────────────────────────────────────
@dataclass
class ContextSchema:
    llm: Any = None                        # langchain_openai.ChatOpenAI instance
    model_config: Any = None               # ModelConfig from sovereign.models
    agent_config: Any = None               # AgentConfig from sovereign.agents
    sandbox_image: str = "sandbox-python:latest"
    sandbox_timeout: int = 60
    sandbox_mem_limit: str = "512m"


# ────────────────────────────────────────────────────────────
# NODE: validate_file
# ────────────────────────────────────────────────────────────
def validate_file(state: OrchestratorState) -> dict:
    """Resolve and validate the target file exists on disk."""
    raw_path = state["file_path"]
    file_path = Path(raw_path).resolve()

    if not file_path.exists():
        raise SystemExit(
            f"[ERROR] File not found: {file_path}\n  Check the path and try again."
        )

    if not file_path.is_file():
        raise SystemExit(f"[ERROR] Path is not a file: {file_path}")

    print(f"  File     : {file_path}")
    print(f"  Directory: {file_path.parent}")

    return {
        "file_path": str(file_path),
        "file_dir": str(file_path.parent),
        "filename": file_path.name,
    }


# ────────────────────────────────────────────────────────────
# NODE: build_prompt
# ────────────────────────────────────────────────────────────
def _file_preview(path: str, max_chars: int = 1500) -> str:
    """Return a compact structural preview of the target file so the LLM
    doesn't have to guess column names or schema.

    Supports .xlsx, .csv, and .docx; returns "" if reading fails so the
    pipeline never breaks just because a preview could not be built.
    """
    try:
        file_path = Path(path)
        ext = file_path.suffix.lower()

        if ext == ".xlsx":
            import openpyxl

            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            try:
                parts: list[str] = []
                for ws in wb.worksheets[:3]:
                    rows = list(ws.iter_rows(values_only=True))
                    if not rows:
                        continue
                    parts.append(f"[Sheet '{ws.title}']")
                    parts.append(
                        "Columns: " + ", ".join(str(h) for h in rows[0])
                    )
                    for row in rows[1:4]:
                        parts.append(" | ".join(str(c) for c in row))
            finally:
                wb.close()
            return "\n".join(parts)[:max_chars]

        if ext == ".csv":
            import csv

            with open(
                file_path, newline="", encoding="utf-8", errors="replace"
            ) as fh:
                rows = list(csv.reader(fh))[:5]
            if not rows:
                return ""
            headings = ["Columns: " + ", ".join(rows[0])]
            headings += [" | ".join(r) for r in rows[1:]]
            return "\n".join(headings)[:max_chars]

        if ext == ".docx":
            from docx import Document

            doc = Document(str(file_path))
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paras[:12])[:max_chars]

    except Exception:
        pass
    return ""


def build_prompt(state: OrchestratorState, runtime: Runtime[ContextSchema]) -> dict:
    """Construct the initial chat messages for the LLM.

    Reads the system prompt from the agent config (YAML-based).
    For registry agents, builds a prompt that outputs YAML file operations.
    On retries this node is skipped; the retry logic appends
    feedback directly to the messages list instead.
    """
    filename = state.get("filename", "")
    user_prompt = state["user_prompt"]
    agent_config = runtime.context.agent_config

    # Give the model the actual file structure instead of letting it guess
    # column/sheet names (a common source of sandbox exit-1 failures).
    preview = _file_preview(state.get("file_path", ""))
    if preview:
        user_prompt = (
            f"{user_prompt}\n\n"
            f"FILE PREVIEW (first rows/columns of `{filename}`):\n{preview}"
        )

    if agent_config:
        system_prompt = agent_config.system_prompt
    else:
        # Fallback for when no agent is configured
        system_prompt = (
            "You are a Python agent running in a sandboxed Docker container.\n"
            "Output ONLY executable Python code. No markdown fences, no explanations.\n"
            "Use sys.exit(0) on success."
        )

    # Substitute {filename} into the system prompt if present
    system_prompt = system_prompt.replace("{filename}", filename)

    messages = [
        SystemMessage(content=system_prompt),
        HumanMessage(content=user_prompt),
    ]

    return {"messages": messages}


# ────────────────────────────────────────────────────────────
# NODE: call_llm
# ────────────────────────────────────────────────────────────
def call_llm(state: OrchestratorState, runtime: Runtime[ContextSchema]) -> dict:
    """Invoke the LLM via the LangChain ChatOpenAI client.

    Supports multimodal messages when the agent requires vision
    and an image_path is provided in state.
    """
    llm = runtime.context.llm
    messages = list(state["messages"])

    # Check if this agent requires vision and image is available
    agent_config = runtime.context.agent_config
    image_path = state.get("image_path")

    if image_path and agent_config and getattr(agent_config, "requires_vision", False):
        from .vision import build_multimodal_message, is_image_file

        if is_image_file(image_path):
            # Find the last HumanMessage and replace with multimodal version
            for i in range(len(messages) - 1, -1, -1):
                if isinstance(messages[i], HumanMessage):
                    text = messages[i].content if isinstance(messages[i].content, str) else ""
                    messages[i] = build_multimodal_message(text, image_path)
                    print(f"  Attached image: {image_path}")
                    break

    print("  Contacting LLM...")
    response = llm.invoke(messages)

    content = response.content
    if not content:
        raise SystemExit("[ERROR] LLM returned empty content")

    print(f"  Received {len(content)} chars from model")

    return {
        "llm_response": content,
        "messages": [AIMessage(content=content)],
    }


# ────────────────────────────────────────────────────────────
# NODE: extract_code
# ────────────────────────────────────────────────────────────
def extract_code(state: OrchestratorState) -> dict:
    """Extract executable Python code from the LLM response.

    Tries fenced code blocks first, then falls back to raw text
    with a sanity check.  On failure, stores feedback so the retry
    loop can pass it back to the LLM.
    """
    response = state["llm_response"]
    attempt = state.get("extract_attempts", 0) + 1

    # Priority 1: ```python ... ``` block
    m = re.search(r"```python\s*\n(.*?)```", response, re.DOTALL)
    if m:
        code = m.group(1).strip()
        print(f"  Extracted {len(code)} chars of Python (fenced python block)")
        return {"code": code, "extract_attempts": attempt, "error_feedback": ""}

    # Priority 2: ``` ... ``` block (any language tag)
    m = re.search(r"```\w*\s*\n(.*?)```", response, re.DOTALL)
    if m:
        code = m.group(1).strip()
        print(f"  Extracted {len(code)} chars of Python (fenced block)")
        return {"code": code, "extract_attempts": attempt, "error_feedback": ""}

    # Priority 3: use the entire response as code
    code = response.strip()
    keywords = ("import ", "def ", "print(", "pd.", "open(", "sys.")
    if code and any(kw in code for kw in keywords):
        print(f"  Extracted {len(code)} chars of Python (raw response)")
        return {"code": code, "extract_attempts": attempt, "error_feedback": ""}

    # Extraction failed — return feedback for the retry loop
    feedback = (
        f"Could not extract valid Python code from your response.\n"
        f"Raw response (first 500 chars):\n{response[:500]}"
    )
    print(
        f"  [WARN] Code extraction failed (attempt {attempt})",
        file=sys.stderr,
    )
    return {"code": "", "extract_attempts": attempt, "error_feedback": feedback}


# ────────────────────────────────────────────────────────────
# NODE: security_scan
# ────────────────────────────────────────────────────────────
def security_scan(state: OrchestratorState, runtime: Runtime[ContextSchema]) -> dict:
    """Run AST + Regex security scan on extracted code.

    Returns the scan result and redacted code.
    """
    from .security.scanner import scan_code

    code = state.get("code", "")
    if not code:
        return {"security_result": {"passed": False, "issues": [], "summary": "No code to scan."}}

    agent_config = runtime.context.agent_config
    agent_security = agent_config.security if agent_config else None

    result = scan_code(code, agent_security)

    print(f"  Security: {result.summary}")

    return {
        "security_result": result.to_dict(),
        "code": result.redacted_code if not result.passed else code,
    }


# ────────────────────────────────────────────────────────────
# NODE: confirm
# ────────────────────────────────────────────────────────────
def confirm(state: OrchestratorState) -> dict:
    """Prompt the user to approve or reject the generated code.

    If state["auto"] is True the code is approved automatically.
    """
    code = state.get("code", "")

    if not code:
        return {"approved": False}

    if state.get("auto"):
        print("  Auto mode — skipping confirmation.")
        return {"approved": True}

    print()
    print("─── Generated Code ───────────────────────────────")
    print(code)
    print("─── End Code ─────────────────────────────────────")
    print()

    try:
        answer = input("Execute this code? [y/N]: ").strip().lower()
    except (EOFError, KeyboardInterrupt):
        print("\nAborted by user.")
        return {"approved": False}

    return {"approved": answer in ("y", "yes")}


# ────────────────────────────────────────────────────────────
# NODE: run_sandbox
# ────────────────────────────────────────────────────────────
def _decode_logs(raw) -> str:
    """Decode docker log bytes into a UTF-8/ASCII-safe string."""
    if not raw:
        return ""
    if isinstance(raw, bytes):
        return raw.decode("utf-8", errors="replace")
    return str(raw)


def run_sandbox(
    state: OrchestratorState, runtime: Runtime[ContextSchema]
) -> dict:
    """Write generated code to a temp file, run it in an isolated
    Docker container, and return the results.

    The container is run detached and removed manually so that the
    real stdout/stderr can be captured *before* cleanup.  (The Docker
    SDK's auto-remove deletes the container before ``ContainerError``
    is built, which turned every sandbox failure into an empty ``b''``
    and hid the actual Python traceback.)
    """
    code = state["code"]
    mount_dir = Path(state["file_dir"])
    agent_config = runtime.context.agent_config

    # Use agent-specific sandbox image if available
    sandbox_image = (
        agent_config.sandbox_image
        if agent_config and agent_config.sandbox_image
        else runtime.context.sandbox_image
    )
    sandbox_timeout = runtime.context.sandbox_timeout
    sandbox_mem_limit = runtime.context.sandbox_mem_limit

    attempt = state.get("sandbox_attempts", 0) + 1
    print(
        f"  Running in sandbox (attempt {attempt}, timeout={sandbox_timeout}s)..."
    )

    try:
        client = docker.from_env()
        client.ping()
    except docker.errors.DockerException as exc:
        return {
            "sandbox_attempts": attempt,
            "sandbox_result": {
                "stdout": "",
                "stderr": (
                    f"Cannot connect to Docker daemon: {exc}\n"
                    "  Make sure Docker Desktop is running."
                ),
                "exit_code": -1,
                "image": sandbox_image,
            },
        }

    script_path = mount_dir / ".__orchestrator_sandbox_script__.py"
    container = None
    exit_code = -1
    stdout = ""
    stderr = ""
    timed_out = False

    try:
        script_path.write_text(code, encoding="utf-8")

        container = client.containers.run(
            image=sandbox_image,
            command=[
                "python",
                "/workspace/.__orchestrator_sandbox_script__.py",
            ],
            volumes={
                str(mount_dir): {"bind": "/workspace", "mode": "rw"}
            },
            network_disabled=True,
            detach=True,
            mem_limit=sandbox_mem_limit,
            cpu_period=100_000,
            cpu_quota=50_000,
        )

        # Wait for the script to finish (wires SANDBOX_TIMEOUT, which was
        # previously never passed to the container run).
        try:
            wait_result = container.wait(timeout=sandbox_timeout)
        except Exception as exc:
            timed_out = True
            exit_code = -1
            stderr = (
                f"Sandbox timed out after {sandbox_timeout}s — "
                f"the script did not exit. {exc}"
            )
            try:
                container.kill()
            except Exception:
                pass
        else:
            exit_code = int(wait_result.get("StatusCode", -1))

        # Fetch logs BEFORE removal so failures keep their tracebacks.
        stdout = _decode_logs(
            container.logs(stdout=True, stderr=False, timestamps=False)
        )
        stderr = (
            _decode_logs(
                container.logs(stdout=False, stderr=True, timestamps=False)
            )
            if not timed_out
            else stderr
        )

        # Distinguish a script that ran cleanly from one that failed.
        if not timed_out:
            stderr = stderr or (
                f"Process exited with status {exit_code} and produced no "
                "error output."
                if exit_code != 0
                else ""
            )

        return {
            "sandbox_attempts": attempt,
            "sandbox_result": {
                "stdout": stdout,
                "stderr": stderr,
                "exit_code": exit_code,
                "image": sandbox_image,
            },
        }

    except docker.errors.ImageNotFound as exc:
        return {
            "sandbox_attempts": attempt,
            "sandbox_result": {
                "stdout": "",
                "stderr": (
                    f"Sandbox image not found: {exc}\n"
                    "  Build it once: docker build -t sandbox-python "
                    "-f Dockerfile.sandbox ."
                ),
                "exit_code": -1,
                "image": sandbox_image,
            },
        }

    except Exception as exc:
        return {
            "sandbox_attempts": attempt,
            "sandbox_result": {
                "stdout": "",
                "stderr": f"Sandbox error: {exc}",
                "exit_code": -1,
                "image": sandbox_image,
            },
        }

    finally:
        if container is not None:
            try:
                container.remove(force=True)
            except Exception:
                pass
        try:
            script_path.unlink(missing_ok=True)
        except OSError:
            pass


# ────────────────────────────────────────────────────────────
# NODE: write_model_config
# ────────────────────────────────────────────────────────────
def write_model_config(
    state: OrchestratorState, runtime: Runtime[ContextSchema]
) -> dict:
    """Write model configuration YAML files to the models/ directory.

    Used by the Model Registry agent to add/replace/remove models.
    The LLM response should contain structured YAML content that
    this node parses and writes to the appropriate file.
    """
    from .models import ModelRegistry

    response = state.get("llm_response", "")
    models_dir = Path(__file__).parent / "models"
    models_dir.mkdir(exist_ok=True)

    # Parse the LLM response to extract YAML blocks
    # The registry agent outputs either:
    # 1. A single YAML block (add/replace)
    # 2. A list of operations
    written_files = []

    # Try to find YAML blocks in the response
    yaml_blocks = re.findall(r"```(?:yaml)?\s*\n(.*?)```", response, re.DOTALL)

    if not yaml_blocks:
        # Try parsing the entire response as YAML
        yaml_blocks = [response.strip()]

    import yaml

    for block in yaml_blocks:
        try:
            data = yaml.safe_load(block)
            if not isinstance(data, dict):
                continue

            name = data.get("name", "")
            if not name:
                continue

            # Create filename from model name
            slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
            if not slug:
                continue

            filepath = models_dir / f"{slug}.yaml"

            # Check if this is a delete operation
            if data.get("_delete"):
                if filepath.exists():
                    filepath.unlink()
                    print(f"  Deleted model config: {filepath.name}")
                continue

            # Write the YAML file
            with open(filepath, "w", encoding="utf-8") as fh:
                yaml.dump(data, fh, default_flow_style=False, allow_unicode=True)
            written_files.append(filepath.name)
            print(f"  Wrote model config: {filepath.name}")

        except yaml.YAMLError as exc:
            print(f"  [WARN] Failed to parse YAML block: {exc}", file=sys.stderr)
            continue

    # Reload the model registry
    try:
        registry = ModelRegistry(models_dir)
        registry.reload()
        print(f"  Model registry reloaded — {len(registry.list_models())} models")
    except Exception as exc:
        print(f"  [WARN] Could not reload model registry: {exc}", file=sys.stderr)

    return {
        "sandbox_result": {
            "stdout": f"Model config(s) written: {', '.join(written_files)}" if written_files else "No model configs were written.",
            "stderr": "",
            "exit_code": 0 if written_files else 1,
        }
    }
